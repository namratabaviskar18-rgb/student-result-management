"""
====================================================
 STUDENT RESULT MANAGEMENT SYSTEM
 A Menu-Driven Command-Line Python Project
 Stores data in an Excel file using openpyxl
====================================================
"""

import os
from openpyxl import Workbook, load_workbook

FILE_NAME = "student_results.xlsx"

HEADERS = [
    "Roll No", "Name", "Class",
    "Subject 1", "Subject 2", "Subject 3", "Subject 4", "Subject 5",
    "Total", "Percentage", "Grade", "Status"
]


# ---------------------------------------------------------
# EXCEL SETUP
# ---------------------------------------------------------
def initialize_excel():
    """Create the Excel file with headers if it doesn't already exist."""
    if not os.path.exists(FILE_NAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        ws.append(HEADERS)
        wb.save(FILE_NAME)


# ---------------------------------------------------------
# CALCULATIONS
# ---------------------------------------------------------
def calculate_result(marks):
    """
    Takes a list of 5 subject marks.
    Returns total, percentage, grade, and pass/fail status.
    """
    total = sum(marks)
    percentage = total / len(marks)

    if percentage >= 90:
        grade = "A+"
    elif percentage >= 80:
        grade = "A"
    elif percentage >= 70:
        grade = "B"
    elif percentage >= 60:
        grade = "C"
    elif percentage >= 40:
        grade = "D"
    else:
        grade = "F"

    # Fail if overall percentage is below 40, OR any single subject is below 33 (typical rule)
    status = "PASS"
    if percentage < 40 or any(m < 33 for m in marks):
        status = "FAIL"

    return total, round(percentage, 2), grade, status


# ---------------------------------------------------------
# INPUT HELPERS
# ---------------------------------------------------------
def get_marks_for_subjects(num_subjects=5, max_marks=100):
    """Loop through subjects and collect valid marks."""
    marks = []
    for i in range(1, num_subjects + 1):
        while True:
            try:
                mark = float(input(f"Enter marks for Subject {i} (out of {max_marks}): "))
                if 0 <= mark <= max_marks:
                    marks.append(mark)
                    break
                else:
                    print(f"⚠️  Please enter a value between 0 and {max_marks}.")
            except ValueError:
                print("⚠️  Invalid input. Please enter a number.")
    return marks


# ---------------------------------------------------------
# 1. ADD STUDENT RESULT
# ---------------------------------------------------------
def add_student():
    print("\n--- Add Student Result ---")

    roll_no = input("Enter Roll No: ").strip()

    # Check for duplicate roll number
    wb = load_workbook(FILE_NAME)
    ws = wb["Results"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None and str(row[0]) == roll_no:
            print(f"⚠️  A student with Roll No {roll_no} already exists. Record not added.")
            return

    name = input("Enter Student Name: ").strip()
    student_class = input("Enter Course/Class: ").strip()

    marks = get_marks_for_subjects()
    total, percentage, grade, status = calculate_result(marks)

    row_data = [roll_no, name, student_class] + marks + [total, percentage, grade, status]
    ws.append(row_data)
    wb.save(FILE_NAME)

    print("\n✅ Student result saved successfully!")
    print(f"   Total: {total}  |  Percentage: {percentage}%  |  Grade: {grade}  |  Status: {status}")


# ---------------------------------------------------------
# 2. GET STUDENT RESULT
# ---------------------------------------------------------
def get_result():
    print("\n--- Get Student Result ---")
    roll_no = input("Enter Roll No to search: ").strip()

    wb = load_workbook(FILE_NAME)
    ws = wb["Results"]

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None and str(row[0]) == roll_no:
            print("\n--------------------------------")
            print("Student Result")
            print("--------------------------------")
            print(f"Roll No     : {row[0]}")
            print(f"Name        : {row[1]}")
            print(f"Class       : {row[2]}")
            print(f"Total       : {row[8]}")
            print(f"Percentage  : {row[9]:.2f}%")
            print(f"Grade       : {row[10]}")
            print(f"Status      : {row[11]}")
            print("--------------------------------")
            return

    print(f"❌ No student found with Roll No {roll_no}.")


# ---------------------------------------------------------
# 3. SHOW ALL STUDENT DATA
# ---------------------------------------------------------
def show_all_data():
    print("\n--- All Student Records ---\n")

    wb = load_workbook(FILE_NAME)
    ws = wb["Results"]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if not rows:
        print("No records found. Please add a student first.")
        return

    header_fmt = f"{'Roll No':<10}{'Name':<15}{'Class':<10}{'Total':<10}{'Percentage':<13}{'Grade':<8}{'Status':<8}"
    print(header_fmt)
    print("-" * len(header_fmt))

    for row in rows:
        if row[0] is None:
            continue
        roll_no, name, student_class = row[0], row[1], row[2]
        total, percentage, grade, status = row[8], row[9], row[10], row[11]
        print(f"{str(roll_no):<10}{str(name):<15}{str(student_class):<10}"
              f"{str(total):<10}{f'{percentage:.2f}':<13}{str(grade):<8}{str(status):<8}")


# ---------------------------------------------------------
# MENU
# ---------------------------------------------------------
def menu():
    initialize_excel()

    while True:
        print("\n========================================")
        print("      STUDENT RESULT MANAGEMENT")
        print("========================================")
        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")

        choice = input("\nEnter your choice: ").strip()

        if choice == "1":
            add_student()
        elif choice == "2":
            get_result()
        elif choice == "3":
            show_all_data()
        elif choice == "4":
            print("\nExiting Student Result Management System. Goodbye! 👋")
            break
        else:
            print("⚠️  Invalid choice. Please enter a number between 1 and 4.")


# ---------------------------------------------------------
# ENTRY POINT
# ---------------------------------------------------------
if __name__ == "__main__":
    menu()
